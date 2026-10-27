import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize presentation with 16:9 Widescreen dimensions
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette Definitions (Dark Executive Tech Theme)
BG_COLOR = RGBColor(11, 17, 32)       # #0B1120 Deep Dark Slate
CARD_BG = RGBColor(22, 32, 54)        # #162036 Card Navy
CARD_BORDER = RGBColor(51, 65, 85)    # #334155 Slate Border
ACCENT_BLUE = RGBColor(56, 189, 248)  # #38BDF8 Cyan/Sky Blue
ACCENT_ORANGE = RGBColor(251, 146, 60)# #FB923C Coral Orange
ACCENT_GREEN = RGBColor(52, 211, 153) # #34D399 Emerald Green
ACCENT_PURPLE = RGBColor(167, 139, 250)# #A78BFA Purple
ACCENT_ROSE = RGBColor(244, 63, 94)   # #F43F5E Rose Red
ACCENT_AMBER = RGBColor(251, 191, 36) # #FBBF24 Amber Gold
TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC Heading White
TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 Muted Slate
TEXT_BODY = RGBColor(226, 232, 240)   # #E2E8F0 Body Text White

blank_layout = prs.slide_layouts[6] # Blank slide

def set_slide_background(slide):
    """Sets a solid premium dark background on the slide."""
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background() # No border
    return bg

def add_header(slide, title_text, category="GLOBAL MOBILE PROCESSOR SUPPLY CHAIN ANALYSIS"):
    """Adds a consistent, sleek executive header banner to the slide."""
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category.upper()
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_BLUE

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.7))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = CARD_BORDER
    line.line.fill.background()

def add_card(slide, left, top, width, height, title, items, accent_color=ACCENT_BLUE, title_icon=""):
    """Adds an executive card container with bullet points and bold highlights."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1)

    # Accent top trim
    trim = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.15), top, width - Inches(0.3), Inches(0.04))
    trim.fill.solid()
    trim.fill.fore_color.rgb = accent_color
    trim.line.fill.background()

    # Content
    tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), width - Inches(0.5), height - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_title = tf.paragraphs[0]
    p_title.text = f"{title_icon} {title}".strip()
    p_title.font.size = Pt(15)
    p_title.font.bold = True
    p_title.font.color.rgb = accent_color
    p_title.space_after = Pt(8)

    for item in items:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        if isinstance(item, tuple):
            label, desc = item
            run_lbl = p.add_run()
            run_lbl.text = f"• {label}: "
            run_lbl.font.bold = True
            run_lbl.font.size = Pt(11)
            run_lbl.font.color.rgb = TEXT_WHITE
            run_desc = p.add_run()
            run_desc.text = desc
            run_desc.font.size = Pt(11)
            run_desc.font.color.rgb = TEXT_MUTED
        else:
            run = p.add_run()
            run.text = f"• {item}"
            run.font.size = Pt(11)
            run.font.color.rgb = TEXT_BODY

def add_metric_box(slide, left, top, width, height, val, label, subtext, color=ACCENT_BLUE):
    """Adds a high-impact KPI statistic block."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = CARD_BG
    box.line.color.rgb = CARD_BORDER
    box.line.width = Pt(1)

    tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.12), width - Inches(0.3), height - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p1 = tf.paragraphs[0]
    p1.text = val
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = color
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(3)

    p3 = tf.add_paragraph()
    p3.text = subtext
    p3.font.size = Pt(9.5)
    p3.font.color.rgb = TEXT_MUTED
    p3.alignment = PP_ALIGN.CENTER

print("Helper functions initialized.")
